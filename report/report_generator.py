"""
報表生成模組
將測試結果匯出為 CSV 或 Excel 格式
"""

import csv
import time
from pathlib import Path
from typing import List

from logic.hall_analyzer import HallAnalysisResult, TestResult as HallTestResult
from logic.encoder_analyzer import EncoderAnalysisResult, TestResult as EncTestResult


class ReportGenerator:
    """
    測試報表生成器
    支援 CSV 與 Excel (xlsx) 格式輸出
    """

    MAX_RECORDS = 10000  # 最多保留筆數，避免記憶體過大

    def __init__(self):
        self._hall_records: List[HallAnalysisResult] = []
        self._enc_records: List[EncoderAnalysisResult] = []

    def add_record(self, hall_result: HallAnalysisResult, enc_result: EncoderAnalysisResult):
        """
        新增一筆測試記錄
        Args:
            hall_result: Hall Sensor 分析結果
            enc_result: Encoder 分析結果
        """
        if len(self._hall_records) < self.MAX_RECORDS:
            self._hall_records.append(hall_result)
        if len(self._enc_records) < self.MAX_RECORDS:
            self._enc_records.append(enc_result)

    def clear(self):
        """清除所有記錄"""
        self._hall_records.clear()
        self._enc_records.clear()

    def get_record_count(self) -> int:
        """取得記錄筆數"""
        return len(self._hall_records)

    # ─── CSV 匯出 ──────────────────────────────────────────────────────────────

    def export_csv(
        self,
        filepath: str,
        hall_history: List[HallAnalysisResult] = None,
        enc_history: List[EncoderAnalysisResult] = None
    ):
        """
        匯出 CSV 報表（Hall 與 Encoder 各一個檔案）
        Args:
            filepath: 輸出路徑（不含副檔名）
            hall_history: Hall 歷史記錄（None 則使用內部記錄）
            enc_history: Encoder 歷史記錄（None 則使用內部記錄）
        """
        hall_data = hall_history or self._hall_records
        enc_data = enc_history or self._enc_records

        base = Path(filepath).with_suffix("")
        hall_path = str(base) + "_hall.csv"
        enc_path = str(base) + "_encoder.csv"

        self._write_hall_csv(hall_path, hall_data)
        self._write_encoder_csv(enc_path, enc_data)

        print(f"[Report] Hall CSV 已儲存: {hall_path}")
        print(f"[Report] Encoder CSV 已儲存: {enc_path}")

    def _write_hall_csv(self, filepath: str, records: List[HallAnalysisResult]):
        """寫入 Hall Sensor CSV"""
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)

            # 標頭
            writer.writerow([
                "時間戳記",
                "Hall U 電壓(V)", "Hall U 準位", "Hall U DI", "Hall U 一致", "Hall U 結果",
                "Hall V 電壓(V)", "Hall V 準位", "Hall V DI", "Hall V 一致", "Hall V 結果",
                "Hall W 電壓(V)", "Hall W 準位", "Hall W DI", "Hall W 一致", "Hall W 結果",
                "整體結果",
            ])

            for rec in records:
                row = [time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(rec.timestamp))]
                for phase in ["U", "V", "W"]:
                    ch = rec.channels.get(phase)
                    if ch:
                        row += [
                            f"{ch.voltage:.4f}",
                            ch.voltage_level.value,
                            "H" if ch.di_state else "L",
                            "✓" if ch.consistency else "✗",
                            ch.overall_result.value,
                        ]
                    else:
                        row += ["---", "---", "---", "---", "---"]
                row.append(rec.overall_result.value)
                writer.writerow(row)

    def _write_encoder_csv(self, filepath: str, records: List[EncoderAnalysisResult]):
        """寫入 Encoder CSV"""
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)

            # 標頭
            writer.writerow([
                "時間戳記",
                "Enc A 電壓(V)", "Enc A 準位", "Enc A DI", "Enc A 一致", "Enc A 結果",
                "Enc B 電壓(V)", "Enc B 準位", "Enc B DI", "Enc B 一致", "Enc B 結果",
                "計數(pulses)", "方向", "轉速(RPM)", "位置(度)",
                "整體結果",
            ])

            for rec in records:
                row = [time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(rec.timestamp))]
                for ch_name in ["A", "B"]:
                    ch = rec.channels.get(ch_name)
                    if ch:
                        row += [
                            f"{ch.voltage:.4f}",
                            ch.voltage_level.value,
                            "H" if ch.di_state else "L",
                            "✓" if ch.consistency else "✗",
                            ch.overall_result.value,
                        ]
                    else:
                        row += ["---", "---", "---", "---", "---"]
                row += [
                    rec.count,
                    rec.direction.value,
                    f"{abs(rec.rpm):.2f}",
                    f"{rec.position_deg:.2f}",
                    rec.overall_result.value,
                ]
                writer.writerow(row)

    # ─── Excel 匯出 ────────────────────────────────────────────────────────────

    def export_excel(
        self,
        filepath: str,
        hall_history: List[HallAnalysisResult] = None,
        enc_history: List[EncoderAnalysisResult] = None
    ):
        """
        匯出 Excel 報表（多工作表）
        Args:
            filepath: 輸出路徑（含 .xlsx 副檔名）
            hall_history: Hall 歷史記錄
            enc_history: Encoder 歷史記錄
        """
        try:
            import openpyxl
            from openpyxl.styles import (
                PatternFill, Font, Alignment, Border, Side
            )
        except ImportError:
            raise ImportError("請安裝 openpyxl: pip install openpyxl")

        hall_data = hall_history or self._hall_records
        enc_data = enc_history or self._enc_records

        wb = openpyxl.Workbook()

        # ── 封面摘要 ──────────────────────────────────────────────────────────
        ws_summary = wb.active
        ws_summary.title = "測試摘要"
        self._write_summary_sheet(ws_summary, hall_data, enc_data, openpyxl)

        # ── Hall Sensor 工作表 ────────────────────────────────────────────────
        ws_hall = wb.create_sheet("Hall Sensor")
        self._write_hall_sheet(ws_hall, hall_data, openpyxl)

        # ── Encoder 工作表 ────────────────────────────────────────────────────
        ws_enc = wb.create_sheet("Encoder")
        self._write_encoder_sheet(ws_enc, enc_data, openpyxl)

        wb.save(filepath)
        print(f"[Report] Excel 報表已儲存: {filepath}")

    def _write_summary_sheet(self, ws, hall_data, enc_data, openpyxl):
        """寫入摘要工作表"""
        from openpyxl.styles import Font, PatternFill, Alignment

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20

        title_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        title_fill = PatternFill("solid", fgColor="1A3A6B")
        header_font = Font(name="Arial", size=11, bold=True)
        center = Alignment(horizontal="center", vertical="center")

        # 標題
        ws.merge_cells("A1:B1")
        ws["A1"] = "馬達測試系統 - 測試報告"
        ws["A1"].font = title_font
        ws["A1"].fill = title_fill
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 30

        # 基本資訊
        info = [
            ("測試時間", time.strftime("%Y-%m-%d %H:%M:%S")),
            ("Hall 記錄筆數", len(hall_data)),
            ("Encoder 記錄筆數", len(enc_data)),
        ]

        # Hall 整體 PASS 率
        if hall_data:
            hall_pass = sum(1 for r in hall_data if r.overall_result.value == "PASS")
            info.append(("Hall PASS 率", f"{hall_pass}/{len(hall_data)} ({hall_pass/len(hall_data)*100:.1f}%)"))
        else:
            info.append(("Hall PASS 率", "無資料"))

        # Encoder 整體 PASS 率
        if enc_data:
            enc_pass = sum(1 for r in enc_data if r.overall_result.value == "PASS")
            info.append(("Encoder PASS 率", f"{enc_pass}/{len(enc_data)} ({enc_pass/len(enc_data)*100:.1f}%)"))
        else:
            info.append(("Encoder PASS 率", "無資料"))

        for row_idx, (key, val) in enumerate(info, start=3):
            ws.cell(row=row_idx, column=1, value=key).font = header_font
            ws.cell(row=row_idx, column=2, value=str(val))

    def _write_hall_sheet(self, ws, records, openpyxl):
        """寫入 Hall Sensor 工作表"""
        from openpyxl.styles import Font, PatternFill, Alignment

        headers = [
            "時間戳記",
            "Hall U 電壓(V)", "Hall U 準位", "Hall U DI", "Hall U 一致", "Hall U 結果",
            "Hall V 電壓(V)", "Hall V 準位", "Hall V DI", "Hall V 一致", "Hall V 結果",
            "Hall W 電壓(V)", "Hall W 準位", "Hall W DI", "Hall W 一致", "Hall W 結果",
            "整體結果",
        ]

        header_fill = PatternFill("solid", fgColor="2D5A8E")
        header_font = Font(name="Arial", bold=True, color="FFFFFF")
        pass_fill = PatternFill("solid", fgColor="1A6B1A")
        fail_fill = PatternFill("solid", fgColor="6B1A1A")
        pass_font = Font(color="00FF00", bold=True)
        fail_font = Font(color="FF4444", bold=True)

        # 表頭
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = 14

        # 資料列
        for row_idx, rec in enumerate(records, start=2):
            col = 1
            ws.cell(row=row_idx, column=col, value=time.strftime("%H:%M:%S", time.localtime(rec.timestamp)))
            col += 1

            for phase in ["U", "V", "W"]:
                ch = rec.channels.get(phase)
                if ch:
                    ws.cell(row=row_idx, column=col, value=round(ch.voltage, 4))
                    ws.cell(row=row_idx, column=col + 1, value=ch.voltage_level.value)
                    ws.cell(row=row_idx, column=col + 2, value="H" if ch.di_state else "L")
                    ws.cell(row=row_idx, column=col + 3, value="✓" if ch.consistency else "✗")
                    result_cell = ws.cell(row=row_idx, column=col + 4, value=ch.overall_result.value)
                    if ch.overall_result.value == "PASS":
                        result_cell.fill = pass_fill
                        result_cell.font = pass_font
                    else:
                        result_cell.fill = fail_fill
                        result_cell.font = fail_font
                col += 5

            overall_cell = ws.cell(row=row_idx, column=col, value=rec.overall_result.value)
            if rec.overall_result.value == "PASS":
                overall_cell.fill = pass_fill
                overall_cell.font = pass_font
            else:
                overall_cell.fill = fail_fill
                overall_cell.font = fail_font

    def _write_encoder_sheet(self, ws, records, openpyxl):
        """寫入 Encoder 工作表"""
        from openpyxl.styles import Font, PatternFill, Alignment

        headers = [
            "時間戳記",
            "Enc A 電壓(V)", "Enc A 準位", "Enc A DI", "Enc A 一致", "Enc A 結果",
            "Enc B 電壓(V)", "Enc B 準位", "Enc B DI", "Enc B 一致", "Enc B 結果",
            "計數(pulses)", "方向", "轉速(RPM)", "位置(度)",
            "整體結果",
        ]

        header_fill = PatternFill("solid", fgColor="2D5A8E")
        header_font = Font(name="Arial", bold=True, color="FFFFFF")
        pass_fill = PatternFill("solid", fgColor="1A6B1A")
        fail_fill = PatternFill("solid", fgColor="6B1A1A")
        pass_font = Font(color="00FF00", bold=True)
        fail_font = Font(color="FF4444", bold=True)

        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = 14

        for row_idx, rec in enumerate(records, start=2):
            col = 1
            ws.cell(row=row_idx, column=col, value=time.strftime("%H:%M:%S", time.localtime(rec.timestamp)))
            col += 1

            for ch_name in ["A", "B"]:
                ch = rec.channels.get(ch_name)
                if ch:
                    ws.cell(row=row_idx, column=col, value=round(ch.voltage, 4))
                    ws.cell(row=row_idx, column=col + 1, value=ch.voltage_level.value)
                    ws.cell(row=row_idx, column=col + 2, value="H" if ch.di_state else "L")
                    ws.cell(row=row_idx, column=col + 3, value="✓" if ch.consistency else "✗")
                    result_cell = ws.cell(row=row_idx, column=col + 4, value=ch.overall_result.value)
                    if ch.overall_result.value == "PASS":
                        result_cell.fill = pass_fill
                        result_cell.font = pass_font
                    else:
                        result_cell.fill = fail_fill
                        result_cell.font = fail_font
                col += 5

            ws.cell(row=row_idx, column=col, value=rec.count)
            ws.cell(row=row_idx, column=col + 1, value=rec.direction.value)
            ws.cell(row=row_idx, column=col + 2, value=round(abs(rec.rpm), 2))
            ws.cell(row=row_idx, column=col + 3, value=round(rec.position_deg, 2))

            overall_cell = ws.cell(row=row_idx, column=col + 4, value=rec.overall_result.value)
            if rec.overall_result.value == "PASS":
                overall_cell.fill = pass_fill
                overall_cell.font = pass_font
            else:
                overall_cell.fill = fail_fill
                overall_cell.font = fail_font
